#!/usr/bin/env python3
# Aim: Download and normalize Global Wind Power Tracker data for DATE Mapper.
# Author: Benjamin Turnbull

"""Download and build the Global Wind Power Tracker GeoJSON layer."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
import re
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TRACKER_PAGE_URL = (
    "https://globalenergymonitor.org/projects/global-wind-power-tracker/"
)
LICENSE_URL = "https://globalenergymonitor.org/creative-commons-license/"
TRACKER_SLUG = "wind-power-tracker"
USER_AGENT = "DATE-Mapper-GWPT-importer/1.0"
DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "data"
    / "High-Resolution"
    / "Infrastructure"
    / "Critical Infrastructure"
    / "Wind Farms"
    / "Global_Wind_Power_Tracker.geojson"
)
DEFAULT_SHEETS = ("Data", "Below Threshold")
USE_CASE_MIN_LENGTH = 100

FORM_ENVIRONMENT_VARIABLES = {
    "name": "GEM_NAME",
    "email": "GEM_EMAIL",
    "organization": "GEM_ORGANIZATION",
    "sector": "GEM_SECTOR",
    "country": "GEM_COUNTRY",
    "use_case": "GEM_USE_CASE",
}

COLUMN_OUTPUT_NAMES = {
    "datelastresearched": "date_last_researched",
    "countryarea": "country",
    "country": "country",
    "projectname": "name",
    "phasename": "phase",
    "projectnameinlocallanguagescript": "project_name_local",
    "othernames": "other_names",
    "capacitymw": "capacity_mw",
    "installationtype": "installation_type",
    "status": "status",
    "startyear": "start_year",
    "retiredyear": "retired_year",
    "operator": "operator",
    "operatornameinlocallanguagescript": "operator_name_local",
    "owner": "owner",
    "ownernameinlocallanguagescript": "owner_name_local",
    "hydrogen": "hydrogen",
    "associatedstorage": "associated_storage",
    "locationaccuracy": "location_accuracy",
    "city": "city",
    "localareatalukcounty": "local_area",
    "majorareaprefecturedistrict": "major_area",
    "stateprovince": "state_province",
    "subregion": "subregion",
    "region": "region",
    "gemlocationid": "gem_location_id",
    "gemphaseid": "gem_phase_id",
    "otheridslocation": "other_location_ids",
    "otheridsunitphase": "other_phase_ids",
    "wikiurl": "wiki_url",
}

PROPERTY_ORDER = (
    "name",
    "phase",
    "country",
    "capacity_mw",
    "installation_type",
    "status",
    "start_year",
    "retired_year",
    "operator",
    "owner",
    "location_accuracy",
    "city",
    "state_province",
    "subregion",
    "region",
    "gem_location_id",
    "gem_phase_id",
    "wiki_url",
    "date_last_researched",
    "project_name_local",
    "other_names",
    "operator_name_local",
    "owner_name_local",
    "hydrogen",
    "associated_storage",
    "local_area",
    "major_area",
    "other_location_ids",
    "other_phase_ids",
)


class DownloadError(RuntimeError):
    """Raised when GEM's download workflow cannot provide a workbook."""


class ScriptSourceParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.sources: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.casefold() != "script":
            return
        attributes = dict(attrs)
        source = attributes.get("src")
        if source:
            self.sources.append(source)


def normalized_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        return int(value) if value.is_integer() else value
    if isinstance(value, str):
        value = value.strip()
        return None if not value or value == "--" else value
    return str(value).strip() or None


def boolean_value(value: Any) -> bool | None:
    value = json_value(value)
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"true", "yes", "y", "1"}:
        return True
    if normalized in {"false", "no", "n", "0"}:
        return False
    return None


def finite_coordinate(value: Any, minimum: float, maximum: float) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(result) or not minimum <= result <= maximum:
        return None
    return result


def source_release(workbook) -> str:
    if "About" not in workbook.sheetnames:
        return "Unknown release"
    for row in workbook["About"].iter_rows(min_row=1, max_row=12, max_col=1, values_only=True):
        text = str(row[0] or "").strip()
        match = re.search(r"Global Wind Power Tracker\s*[-–]\s*(.+)", text, re.I)
        if match:
            return match.group(1).strip()
    return "Unknown release"


def feature_identifier(properties: dict, sheet_name: str, row_number: int) -> str:
    phase_id = properties.get("gem_phase_id")
    if phase_id:
        return str(phase_id)
    identity = "\0".join(
        str(properties.get(key) or "")
        for key in ("gem_location_id", "name", "phase", "country")
    )
    identity = f"{sheet_name}\0{row_number}\0{identity}"
    return f"gwpt-{hashlib.sha256(identity.encode('utf-8')).hexdigest()[:20]}"


def normalized_properties(row_by_header: dict[str, Any], below_threshold: bool) -> dict:
    converted = {}
    for source_name, output_name in COLUMN_OUTPUT_NAMES.items():
        value = row_by_header.get(source_name)
        if output_name in {"hydrogen", "associated_storage"}:
            value = boolean_value(value)
        else:
            value = json_value(value)
        if value is not None:
            converted[output_name] = value

    properties = {
        key: converted[key]
        for key in PROPERTY_ORDER
        if key in converted
    }
    properties["below_threshold"] = below_threshold
    properties["DATASET"] = "WORLD"
    properties["source"] = "Global Energy Monitor - Global Wind Power Tracker"
    return properties


def features_from_sheet(worksheet, sheet_name: str) -> tuple[list[dict], int]:
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = next(rows)
    except StopIteration as exc:
        raise ValueError(f"Workbook sheet {sheet_name!r} is empty") from exc

    normalized_headers = [normalized_header(value) for value in headers]
    required = {"projectname", "latitude", "longitude"}
    missing = sorted(required.difference(normalized_headers))
    if missing:
        raise ValueError(
            f"Workbook sheet {sheet_name!r} is missing required columns: "
            + ", ".join(missing)
        )

    below_threshold = normalized_header(sheet_name) == "belowthreshold"
    features = []
    skipped_coordinates = 0
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None for value in row):
            continue
        row_by_header = dict(zip(normalized_headers, row))
        latitude = finite_coordinate(row_by_header.get("latitude"), -90.0, 90.0)
        longitude = finite_coordinate(row_by_header.get("longitude"), -180.0, 180.0)
        if latitude is None or longitude is None:
            skipped_coordinates += 1
            continue

        properties = normalized_properties(row_by_header, below_threshold)
        if not properties.get("name"):
            continue
        features.append({
            "type": "Feature",
            "id": feature_identifier(properties, sheet_name, row_number),
            "properties": properties,
            "geometry": {
                "type": "Point",
                "coordinates": [longitude, latitude],
            },
        })
    return features, skipped_coordinates


def workbook_bytes_from_download(payload: bytes) -> bytes:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            names = archive.namelist()
            if "xl/workbook.xml" in names:
                return payload
            workbooks = [name for name in names if name.casefold().endswith(".xlsx")]
            if not workbooks:
                raise ValueError("Downloaded ZIP does not contain an XLSX workbook")
            preferred = [name for name in workbooks if "wind" in name.casefold()]
            selected = preferred[0] if len(preferred) == 1 else workbooks[0]
            return archive.read(selected)
    except zipfile.BadZipFile as exc:
        raise ValueError("Downloaded file is not a valid XLSX or ZIP archive") from exc


def build_wind_farm_geojson(
    workbook_payload: bytes,
    include_below_threshold: bool = True,
) -> dict:
    workbook_payload = workbook_bytes_from_download(workbook_payload)
    try:
        workbook = load_workbook(
            io.BytesIO(workbook_payload),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(f"Unable to read the GWPT workbook: {exc}") from exc

    selected_sheets = [DEFAULT_SHEETS[0]]
    if include_below_threshold:
        selected_sheets.append(DEFAULT_SHEETS[1])
    missing_sheets = [name for name in selected_sheets if name not in workbook.sheetnames]
    if missing_sheets:
        raise ValueError(
            "GWPT workbook is missing required sheets: " + ", ".join(missing_sheets)
        )

    features = []
    skipped_coordinates = 0
    sheet_counts = {}
    for sheet_name in selected_sheets:
        sheet_features, sheet_skipped = features_from_sheet(
            workbook[sheet_name],
            sheet_name,
        )
        features.extend(sheet_features)
        skipped_coordinates += sheet_skipped
        sheet_counts[sheet_name] = len(sheet_features)

    feature_ids = [feature["id"] for feature in features]
    if len(feature_ids) != len(set(feature_ids)):
        raise ValueError("GWPT workbook contains duplicate GEM phase identifiers")
    if not features:
        raise ValueError("GWPT workbook contains no wind farms with valid coordinates")

    features.sort(key=lambda feature: (
        str(feature["properties"].get("country") or "").casefold(),
        str(feature["properties"].get("name") or "").casefold(),
        str(feature["id"]),
    ))
    longitudes = [feature["geometry"]["coordinates"][0] for feature in features]
    latitudes = [feature["geometry"]["coordinates"][1] for feature in features]
    release = source_release(workbook)
    workbook.close()

    return {
        "type": "FeatureCollection",
        "name": "Global Wind Power Tracker",
        "bbox": [
            min(longitudes),
            min(latitudes),
            max(longitudes),
            max(latitudes),
        ],
        "source": {
            "name": "Global Energy Monitor",
            "dataset": "Global Wind Power Tracker",
            "release": release,
            "url": TRACKER_PAGE_URL,
            "license": "Creative Commons Attribution 4.0 International",
            "license_url": LICENSE_URL,
            "citation": (
                "Global Energy Monitor, Global Wind Power Tracker, "
                f"{release} release"
            ),
            "source_sha256": hashlib.sha256(workbook_payload).hexdigest(),
        },
        "import_summary": {
            "feature_count": len(features),
            "skipped_invalid_coordinates": skipped_coordinates,
            "included_sheets": selected_sheets,
            "sheet_feature_counts": sheet_counts,
        },
        "features": features,
    }


def write_geojson(payload: dict, output: Path, pretty: bool = False) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        dir=output.parent,
        prefix=f".{output.name}.",
        suffix=".tmp",
        delete=False,
    ) as temp_file:
        temporary_path = Path(temp_file.name)
        json.dump(
            payload,
            temp_file,
            ensure_ascii=False,
            indent=2 if pretty else None,
            separators=None if pretty else (",", ":"),
        )
        temp_file.write("\n")
    temporary_path.chmod(0o644)
    temporary_path.replace(output)


def write_workbook(payload: bytes, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="wb",
        dir=output.parent,
        prefix=f".{output.name}.",
        suffix=".tmp",
        delete=False,
    ) as temp_file:
        temporary_path = Path(temp_file.name)
        temp_file.write(workbook_bytes_from_download(payload))
    temporary_path.chmod(0o644)
    temporary_path.replace(output)


def fetch_bytes(
    url: str,
    *,
    method: str = "GET",
    headers: dict[str, str] | None = None,
    body: bytes | None = None,
    timeout: int = 120,
) -> bytes:
    request_headers = {"User-Agent": USER_AGENT, **(headers or {})}
    request = urllib.request.Request(
        url,
        data=body,
        headers=request_headers,
        method=method,
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return response.read()
    except urllib.error.HTTPError as exc:
        detail = exc.read(500).decode("utf-8", errors="replace").strip()
        suffix = f": {detail}" if detail else ""
        raise DownloadError(f"HTTP {exc.code} from {url}{suffix}") from exc
    except urllib.error.URLError as exc:
        raise DownloadError(f"Unable to reach {url}: {exc.reason}") from exc


def parse_form_script_url(page_html: str, page_url: str = TRACKER_PAGE_URL) -> str:
    parser = ScriptSourceParser()
    parser.feed(page_html)
    candidates = [source for source in parser.sources if "gem-download-form" in source]
    if not candidates:
        raise DownloadError("GEM tracker page no longer exposes its download-form script")
    return urllib.parse.urljoin(page_url, candidates[-1])


def parse_download_form_config(script_text: str) -> dict[str, str]:
    attributes = {
        "submissions_url": ("_submissionsUrl", "submissions-url"),
        "supabase_key": ("_supabaseKey", "supabase-key"),
        "presign_url": ("_presignUrl", "presign-url"),
    }
    config = {}
    for result_name, (getter_name, attribute_name) in attributes.items():
        pattern = (
            rf"get\s+{re.escape(getter_name)}\s*\(\)\s*\{{\s*return\s+"
            rf"this\.getAttribute\([\"']{re.escape(attribute_name)}[\"']\)"
            rf"\s*\|\|\s*[\"']([^\"']+)[\"']"
        )
        match = re.search(pattern, script_text)
        if not match:
            raise DownloadError(
                f"Unable to discover {attribute_name} from GEM's download form"
            )
        config[result_name] = match.group(1)
    return config


def discover_download_form_config() -> dict[str, str]:
    page_html = fetch_bytes(TRACKER_PAGE_URL).decode("utf-8", errors="replace")
    script_url = parse_form_script_url(page_html)
    script_text = fetch_bytes(script_url).decode("utf-8", errors="replace")
    return parse_download_form_config(script_text)


def parse_json_response(payload: bytes, source: str) -> Any:
    try:
        return json.loads(payload.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise DownloadError(f"Invalid JSON response from {source}") from exc


def download_official_workbook(form_values: dict[str, str], email_opt_in: bool) -> bytes:
    config = discover_download_form_config()
    submission = {
        "name": form_values["name"],
        "email": form_values["email"],
        "organization": form_values["organization"],
        "sector": form_values["sector"],
        "country": form_values.get("country", ""),
        "use_case": form_values["use_case"],
        "license_text": (
            "Creative Commons Attribution 4.0 International (CC BY 4.0) - "
            + LICENSE_URL
        ),
        "email_optin": email_opt_in,
        "request_mode": "slugs",
        "useragent": USER_AGENT,
        "page_url": TRACKER_PAGE_URL,
        "requested_slugs": [TRACKER_SLUG],
    }
    api_headers = {
        "Content-Type": "application/json",
        "apikey": config["supabase_key"],
        "Authorization": f"Bearer {config['supabase_key']}",
    }
    minted_payload = fetch_bytes(
        config["submissions_url"],
        method="POST",
        headers=api_headers,
        body=json.dumps(submission).encode("utf-8"),
    )
    minted = parse_json_response(minted_payload, "GEM submission API")
    capability_token = minted.get("capability_token") if isinstance(minted, dict) else None
    if not capability_token:
        raise DownloadError("GEM submission API did not return a download capability")

    presigned_payload = fetch_bytes(
        config["presign_url"],
        method="POST",
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {capability_token}",
        },
        body=b"",
    )
    presigned = parse_json_response(presigned_payload, "GEM presign API")
    urls = presigned.get("urls") if isinstance(presigned, dict) else None
    if not isinstance(urls, list) or not urls:
        raise DownloadError("GEM presign API returned no download URLs")

    candidates = [
        item
        for item in urls
        if isinstance(item, dict)
        and isinstance(item.get("url"), str)
        and (
            item.get("slug") == TRACKER_SLUG
            or str(item.get("filename") or "").casefold().endswith(".xlsx")
        )
    ]
    if not candidates:
        raise DownloadError("GEM response did not contain a wind tracker workbook")
    return fetch_bytes(candidates[0]["url"])


def prompt_value(current: str | None, prompt: str, required: bool = True) -> str:
    value = str(current or "").strip()
    if value:
        return value
    if sys.stdin.isatty():
        value = input(f"{prompt}: ").strip()
    if required and not value:
        raise ValueError(f"{prompt} is required for the official GEM download form")
    return value


def official_form_values(args: argparse.Namespace) -> dict[str, str]:
    values = {
        key: getattr(args, key) or os.environ.get(environment_name, "")
        for key, environment_name in FORM_ENVIRONMENT_VARIABLES.items()
    }
    values["name"] = prompt_value(values["name"], "Name")
    values["email"] = prompt_value(values["email"], "Email address")
    values["organization"] = prompt_value(values["organization"], "Organization")
    values["sector"] = prompt_value(values["sector"], "Sector")
    values["country"] = prompt_value(values["country"], "Country", required=False)
    values["use_case"] = prompt_value(
        values["use_case"],
        f"Intended use ({USE_CASE_MIN_LENGTH}+ characters)",
    )

    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", values["email"]):
        raise ValueError("Email address is not valid")
    if len(values["use_case"]) < USE_CASE_MIN_LENGTH:
        raise ValueError(
            f"Intended use must contain at least {USE_CASE_MIN_LENGTH} characters"
        )
    return values


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    source_group = parser.add_mutually_exclusive_group()
    source_group.add_argument(
        "--source-file",
        type=Path,
        help="Convert an already-downloaded GEM XLSX file instead of downloading.",
    )
    source_group.add_argument(
        "--download-url",
        help="Download from an existing GEM-issued or direct XLSX URL.",
    )
    parser.add_argument("--name", help="GEM download-form name (or GEM_NAME).")
    parser.add_argument("--email", help="GEM download-form email (or GEM_EMAIL).")
    parser.add_argument(
        "--organization",
        help="GEM download-form organization (or GEM_ORGANIZATION).",
    )
    parser.add_argument("--sector", help="GEM download-form sector (or GEM_SECTOR).")
    parser.add_argument("--country", help="Optional GEM form country (or GEM_COUNTRY).")
    parser.add_argument(
        "--use-case",
        dest="use_case",
        help="100+ character intended use description (or GEM_USE_CASE).",
    )
    parser.add_argument(
        "--accept-license",
        action="store_true",
        help="Confirm acceptance of GEM's CC BY 4.0 licence for an online download.",
    )
    parser.add_argument(
        "--email-opt-in",
        action="store_true",
        help="Opt in to GEM project email updates; disabled by default.",
    )
    parser.add_argument(
        "--exclude-below-threshold",
        action="store_true",
        help="Import only the main Data sheet and omit below-10-MW records.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Destination for application-ready GeoJSON.",
    )
    parser.add_argument(
        "--workbook-output",
        type=Path,
        help="Optionally retain a copy of the downloaded source XLSX.",
    )
    parser.add_argument(
        "--pretty",
        action="store_true",
        help="Indent the GeoJSON. Compact output is substantially smaller.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        if args.source_file:
            workbook_payload = args.source_file.read_bytes()
            source_description = str(args.source_file)
        elif args.download_url:
            workbook_payload = fetch_bytes(args.download_url)
            source_description = args.download_url
        else:
            if not args.accept_license:
                raise ValueError(
                    "--accept-license is required before submitting GEM's download form; "
                    f"review {LICENSE_URL}"
                )
            workbook_payload = download_official_workbook(
                official_form_values(args),
                args.email_opt_in,
            )
            source_description = "the official GEM download form"

        workbook_payload = workbook_bytes_from_download(workbook_payload)
        result = build_wind_farm_geojson(
            workbook_payload,
            include_below_threshold=not args.exclude_below_threshold,
        )
        write_geojson(result, args.output, pretty=args.pretty)
        if args.workbook_output:
            write_workbook(workbook_payload, args.workbook_output)

        summary = result["import_summary"]
        print(
            f"Wrote {summary['feature_count']} wind-farm phases from "
            f"{source_description} to {args.output}"
        )
        if summary["skipped_invalid_coordinates"]:
            print(
                "Skipped "
                f"{summary['skipped_invalid_coordinates']} rows with invalid coordinates"
            )
        return 0
    except (DownloadError, OSError, ValueError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
