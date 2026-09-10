# Aim: Verify Global Wind Power Tracker discovery, conversion, and output generation.
# Author: Benjamin Turnbull

import io
import json
import stat
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from Utilities.update_global_wind_power_tracker import (
    DEFAULT_OUTPUT,
    build_wind_farm_geojson,
    parse_download_form_config,
    parse_form_script_url,
    workbook_bytes_from_download,
    write_geojson,
)


HEADERS = [
    "Date Last Researched",
    "Country/Area",
    "Project Name",
    "Phase Name",
    "Capacity (MW)",
    "Installation Type",
    "Status",
    "Start year",
    "Operator",
    "Owner",
    "Hydrogen",
    "Associated Storage",
    "Latitude",
    "Longitude",
    "Location accuracy",
    "City",
    "State/Province",
    "Subregion",
    "Region",
    "GEM location ID",
    "GEM phase ID",
    "Wiki URL",
]


def sample_row(
    name="Mølle wind farm",
    latitude=55.5,
    longitude=8.25,
    phase_id="G100000000001",
):
    return [
        "2026/01/02",
        "Denmark",
        name,
        1.0,
        120.5,
        "Onshore",
        "operating",
        2024.0,
        "Example Operator",
        "Example Owner [100%]",
        "True",
        None,
        latitude,
        longitude,
        "exact",
        "Esbjerg",
        "Region of Southern Denmark",
        "Northern Europe",
        "Europe",
        "L100000000001",
        phase_id,
        "https://www.gem.wiki/Example_wind_farm",
    ]


def sample_workbook_bytes():
    workbook = Workbook()
    about = workbook.active
    about.title = "About"
    about.append(["Global Wind Power Tracker - February 2099"])
    about.append([
        "Copyright Global Energy Monitor. Distributed under a Creative "
        "Commons Attribution 4.0 International License."
    ])

    data = workbook.create_sheet("Data")
    data.append(HEADERS)
    data.append(sample_row())
    data.append(sample_row(name="Invalid location", latitude=200, phase_id="bad"))

    below_threshold = workbook.create_sheet("Below Threshold")
    below_threshold.append(HEADERS)
    below_threshold.append(sample_row(
        name="Small wind farm",
        latitude=-35.1,
        longitude=117.3,
        phase_id="G100000000002",
    ))

    payload = io.BytesIO()
    workbook.save(payload)
    workbook.close()
    return payload.getvalue()


class UpdateGlobalWindPowerTrackerTests(unittest.TestCase):
    def test_builds_utf8_geojson_from_both_data_sheets(self):
        result = build_wind_farm_geojson(sample_workbook_bytes())

        self.assertEqual(result["source"]["release"], "February 2099")
        self.assertEqual(result["import_summary"]["feature_count"], 2)
        self.assertEqual(result["import_summary"]["skipped_invalid_coordinates"], 1)
        self.assertEqual(result["bbox"], [8.25, -35.1, 117.3, 55.5])

        feature = next(
            item for item in result["features"]
            if item["properties"]["name"] == "Mølle wind farm"
        )
        self.assertEqual(feature["id"], "G100000000001")
        self.assertEqual(feature["properties"]["phase"], 1)
        self.assertEqual(feature["properties"]["capacity_mw"], 120.5)
        self.assertTrue(feature["properties"]["hydrogen"])
        self.assertFalse(feature["properties"]["below_threshold"])
        self.assertEqual(feature["properties"]["DATASET"], "WORLD")
        self.assertEqual(feature["geometry"]["coordinates"], [8.25, 55.5])

    def test_can_exclude_below_threshold_sheet(self):
        result = build_wind_farm_geojson(
            sample_workbook_bytes(),
            include_below_threshold=False,
        )

        self.assertEqual(result["import_summary"]["feature_count"], 1)
        self.assertEqual(result["import_summary"]["included_sheets"], ["Data"])

    def test_extracts_xlsx_from_a_download_zip(self):
        workbook_payload = sample_workbook_bytes()
        archive_payload = io.BytesIO()
        with zipfile.ZipFile(archive_payload, "w") as archive:
            archive.writestr("Global-Wind-Power-Tracker.xlsx", workbook_payload)

        extracted = workbook_bytes_from_download(archive_payload.getvalue())

        self.assertEqual(extracted, workbook_payload)

    def test_parses_current_download_form_configuration_shape(self):
        html = (
            '<script type="module" '
            'src="https://api.example/gem-download-form.bundle.js"></script>'
        )
        self.assertEqual(
            parse_form_script_url(html, "https://globalenergymonitor.example/tracker"),
            "https://api.example/gem-download-form.bundle.js",
        )

        script = (
            'get _submissionsUrl(){return this.getAttribute("submissions-url")||'
            '"https://api.example/submit"}'
            'get _supabaseKey(){return this.getAttribute("supabase-key")||"public-key"}'
            'get _presignUrl(){return this.getAttribute("presign-url")||'
            '"https://api.example/presign"}'
        )
        self.assertEqual(
            parse_download_form_config(script),
            {
                "submissions_url": "https://api.example/submit",
                "supabase_key": "public-key",
                "presign_url": "https://api.example/presign",
            },
        )

    def test_writes_valid_utf8_geojson_atomically(self):
        payload = build_wind_farm_geojson(sample_workbook_bytes())
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "nested" / "wind.geojson"

            write_geojson(payload, output)

            self.assertIn("Mølle wind farm", output.read_text(encoding="utf-8"))
            self.assertEqual(json.loads(output.read_text(encoding="utf-8")), payload)
            self.assertEqual(stat.S_IMODE(output.stat().st_mode), 0o644)

    def test_default_output_uses_critical_infrastructure_structure(self):
        self.assertEqual(DEFAULT_OUTPUT.name, "Global_Wind_Power_Tracker.geojson")
        self.assertEqual(
            DEFAULT_OUTPUT.parts[-5:-1],
            ("High-Resolution", "Infrastructure", "Critical Infrastructure", "Wind Farms"),
        )


if __name__ == "__main__":
    unittest.main()
